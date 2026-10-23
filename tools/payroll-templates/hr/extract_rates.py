"""Extract the local income-tax rates from the Tax Administration's 2026 table into municipal-rates-2026.csv.

The table is "Tablica poreznih stopa godišnjeg poreza na dohodak za 2026. godinu" (Porezna uprava,
from the decisions of the local units published in Narodne novine; state at 1 January 2026).
The download is checked against the SHA-256 recorded below before anything is read.

Usage: python3 -I extract_rates.py <downloaded .xlsx> municipal-rates-2026.csv
"""
import csv
import hashlib
import sys
from decimal import Decimal

import openpyxl

SHA256 = "87733f3eb50a39650813a4ea7a6efbb8c556d4bf9c119ece9ea411f37e6f1592"


def main():
    xlsx, out = sys.argv[1], sys.argv[2]
    with open(xlsx, "rb") as f:
        digest = hashlib.sha256(f.read()).hexdigest()
    if digest != SHA256:
        sys.exit(f"{xlsx}: SHA-256 {digest} is not the 2026 table ({SHA256}); check whether the table was republished")
    sheet = openpyxl.load_workbook(xlsx, read_only=True, data_only=True).active
    header = next(sheet.iter_rows(max_row=1, values_only=True))
    if header[1] != "Šifra grada/općine" or not str(header[5]).startswith("Niža stopa 2026") or not str(header[6]).startswith("Viša stopa 2026"):
        sys.exit(f"unexpected columns {header}")
    rows = []
    for r in sheet.iter_rows(min_row=2, values_only=True):
        if r[1] is None:
            continue
        lower, higher = Decimal(str(r[5])), Decimal(str(r[6]))
        rows.append({"code": int(r[1]), "name": r[2].strip(), "lower": str(lower), "higher": str(higher),
                     "gazette": str(r[4]).strip(), "inForceFrom": str(r[7]).strip()})
    if len(rows) != 556 or len({r["code"] for r in rows}) != 556:
        sys.exit(f"expected 556 distinct local units, found {len(rows)}")
    with open(out, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]), lineterminator="\n")
        w.writeheader()
        w.writerows(sorted(rows, key=lambda r: r["code"]))
    print("wrote", out, "with", len(rows), "local units")


main()
